"""Artifact-backed JSON, Excel, and Markdown experiment reporting."""

from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


REGISTRY_SHEETS = ("PILOT", "MAIN", "ABLATION", "DIAGNOSTIC", "HPARAM")


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def flatten(value: Any, prefix: str = "") -> list[tuple[str, Any]]:
    rows = []
    if isinstance(value, dict):
        for key in sorted(value):
            child = f"{prefix}.{key}" if prefix else str(key)
            rows.extend(flatten(value[key], child))
    elif isinstance(value, list):
        rows.append((prefix, json.dumps(value, sort_keys=True)))
    else:
        rows.append((prefix, value))
    return rows


def read_registry(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = []
    for sheet_name in REGISTRY_SHEETS:
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if not values[0]:
                continue
            row = dict(zip(headers, values, strict=True))
            row["registry_sheet"] = sheet_name
            rows.append(row)
    return rows


def load_result(row: dict[str, Any]) -> tuple[dict[str, Any] | None, str]:
    raw_path = row.get("raw_result_path")
    if not raw_path:
        return None, "not_applicable"
    path = Path(str(raw_path))
    if path.is_dir():
        path = path / "result.json"
    if not path.exists():
        return None, "missing"
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None, "invalid_json"
    if not isinstance(payload, dict):
        return None, "invalid_payload"
    if payload.get("run_id") not in {None, row["run_id"]}:
        return None, "run_id_mismatch"
    if "status" not in payload:
        return None, "missing_status"
    return payload, "loaded"


def render_reports(
    registry_path: Path, output_xlsx: Path, output_markdown: Path
) -> dict[str, Any]:
    registry_rows = read_registry(registry_path)
    resolved = []
    artifacts = []
    metrics = []
    for row in registry_rows:
        result, result_state = load_result(row)
        raw_path = row.get("raw_result_path")
        resolved_raw_path = Path(str(raw_path)) if raw_path else None
        if resolved_raw_path is not None and resolved_raw_path.is_dir():
            resolved_raw_path = resolved_raw_path / "result.json"
        raw_sha = (
            sha256_file(resolved_raw_path)
            if result_state == "loaded" and resolved_raw_path is not None
            else ""
        )
        resolved.append(
            {
                "run_id": row["run_id"],
                "sheet": row["registry_sheet"],
                "todo_id": row.get("todo_id") or "",
                "status": row.get("status") or "",
                "result_state": result_state,
                "result_status": result.get("status", "") if result else "",
                "raw_result_path": str(resolved_raw_path) if resolved_raw_path else "",
                "raw_result_sha256": raw_sha,
                "failure_reason": row.get("failure_reason") or "",
                "notes": row.get("notes") or "",
            }
        )
        if result:
            for key, value in flatten(result.get("metrics", {})):
                metrics.append(
                    {"run_id": row["run_id"], "metric": key, "value": value}
                )
            for label, artifact_path in sorted(result.get("artifacts", {}).items()):
                path = Path(str(artifact_path))
                artifacts.append(
                    {
                        "run_id": row["run_id"],
                        "label": label,
                        "path": str(path),
                        "exists": path.exists(),
                        "sha256": sha256_file(path) if path.is_file() else "",
                    }
                )

    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, data in (
        ("Runs", resolved),
        ("Metrics", metrics),
        ("Artifact_Hashes", artifacts),
    ):
        sheet = workbook.create_sheet(title)
        headers = list(data[0]) if data else ["empty"]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for item in data:
            sheet.append([item[key] for key in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_xlsx)

    lines = [
        "# BADIT-TF Structured Results Summary",
        "",
        f"Generated from `{registry_path}`. Missing results remain explicit.",
        "",
        "| Run | Group | Status | Result | Result SHA256 |",
        "|---|---|---|---|---|",
    ]
    for row in resolved:
        lines.append(
            f"| {row['run_id']} | {row['sheet']} | {row['status']} | "
            f"{row['result_state']} | {row['raw_result_sha256']} |"
        )
    lines.extend(
        [
            "",
            "## P1 recovery decision",
            "",
        ]
    )
    p1_r1 = next(
        (item for item in resolved if item["run_id"] == "p1_r1_confirmation_seed1"),
        None,
    )
    if p1_r1 and p1_r1["result_state"] == "loaded":
        result = json.loads(Path(p1_r1["raw_result_path"]).read_text())
        m = result["metrics"]
        lines.extend(
            [
                f"- Decision: `{result['decision_status']}`.",
                f"- Locked eta: `{result['locked_eta']}`.",
                f"- Spearman: `{m['spearman']}`; task-bootstrap 95% CI: "
                f"`{m['spearman_task_bootstrap_95ci']}`.",
                f"- Records retained: `{m['records']}/{m['expected_records']}`.",
            ]
        )
    else:
        lines.append("- P1-R1 result unavailable.")
    output_markdown.write_text("\n".join(lines) + "\n", encoding="utf-8")
    result = {
        "run_id": "imp6_structured_writers",
        "status": "complete",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_registry": str(registry_path),
        "source_registry_sha256": sha256_file(registry_path),
        "outputs": {
            "xlsx": str(output_xlsx),
            "xlsx_sha256": sha256_file(output_xlsx),
            "markdown": str(output_markdown),
            "markdown_sha256": sha256_file(output_markdown),
        },
        "counts": {
            "registered_runs": len(resolved),
            "loaded_results": sum(
                item["result_state"] == "loaded" for item in resolved
            ),
            "missing_results": sum(
                item["result_state"] == "missing" for item in resolved
            ),
            "metric_rows": len(metrics),
            "artifact_rows": len(artifacts),
        },
        "assertions": {
            "p1_original_retained_failed": any(
                item["run_id"] == "p1_qwen3_4b_objective_seed1"
                and item["status"] == "failed"
                for item in resolved
            ),
            "p1_r1_retained_complete": any(
                item["run_id"] == "p1_r1_confirmation_seed1"
                and item["status"] == "complete"
                for item in resolved
            ),
            "missing_results_explicit": all(
                item["result_state"]
                in {
                    "loaded",
                    "missing",
                    "not_applicable",
                    "invalid_json",
                    "invalid_payload",
                    "run_id_mismatch",
                    "missing_status",
                }
                for item in resolved
            ),
            "xlsx_written": output_xlsx.exists(),
            "markdown_written": output_markdown.exists(),
        },
    }
    if not all(result["assertions"].values()):
        result["status"] = "failed"
    return result
