import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from badit_tf.reporting import render_reports


def test_render_reports_preserves_failed_and_missing(tmp_path: Path) -> None:
    registry = tmp_path / "registry.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    headers = [
        "run_id",
        "status",
        "todo_id",
        "raw_result_path",
        "failure_reason",
        "notes",
    ]
    result_path = tmp_path / "p1.json"
    result_path.write_text(
        json.dumps({"run_id": "p1_qwen3_4b_objective_seed1", "status": "failed"})
    )
    p1_r1_path = tmp_path / "p1r1.json"
    p1_r1_path.write_text(
        json.dumps(
            {
                "run_id": "p1_r1_confirmation_seed1",
                "status": "complete",
                "decision_status": "PASSED_AFTER_PROTOCOL_AMENDMENT",
                "locked_eta": 0.0001,
                "metrics": {
                    "spearman": 1.0,
                    "spearman_task_bootstrap_95ci": [0.9, 1.0],
                    "records": 1,
                    "expected_records": 1,
                },
            }
        )
    )
    for name in ("PILOT", "MAIN", "ABLATION", "DIAGNOSTIC", "HPARAM"):
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        if name == "PILOT":
            sheet.append(
                [
                    "p1_qwen3_4b_objective_seed1",
                    "failed",
                    "P1",
                    str(result_path),
                    "failed",
                    "",
                ]
            )
            sheet.append(
                [
                    "p1_r1_confirmation_seed1",
                    "complete",
                    "P1-R1",
                    str(p1_r1_path),
                    "",
                    "",
                ]
            )
            sheet.append(["p2", "blocked", "P2", str(tmp_path / "missing.json"), "", ""])
    workbook.save(registry)
    xlsx = tmp_path / "out.xlsx"
    markdown = tmp_path / "out.md"
    result = render_reports(registry, xlsx, markdown)
    assert result["status"] == "complete"
    rows = list(load_workbook(xlsx, read_only=True)["Runs"].values)
    assert rows[1][2] == "P1"
    assert "p2" in markdown.read_text()


def test_render_reports_resolves_result_directory(tmp_path: Path) -> None:
    registry = tmp_path / "registry.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    headers = [
        "run_id",
        "status",
        "todo_id",
        "raw_result_path",
        "failure_reason",
        "notes",
    ]
    result_dir = tmp_path / "p3"
    result_dir.mkdir()
    (result_dir / "result.json").write_text(
        json.dumps({"run_id": "p3", "status": "complete", "metrics": {"x": 1}})
    )
    for name in ("PILOT", "MAIN", "ABLATION", "DIAGNOSTIC", "HPARAM"):
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        if name == "PILOT":
            sheet.append(["p3", "complete", "P3", str(result_dir), "", ""])
    workbook.save(registry)
    xlsx = tmp_path / "out.xlsx"
    render_reports(registry, xlsx, tmp_path / "out.md")
    rows = list(load_workbook(xlsx, read_only=True)["Runs"].values)
    assert rows[1][6] == str(result_dir / "result.json")
    assert rows[1][7]


def test_render_reports_preserves_mismatched_result_as_audit_state(tmp_path: Path) -> None:
    registry = tmp_path / "registry.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    headers = [
        "run_id",
        "status",
        "todo_id",
        "raw_result_path",
        "failure_reason",
        "notes",
    ]
    result_path = tmp_path / "wrong.json"
    result_path.write_text(json.dumps({"run_id": "different", "status": "complete"}))
    for name in ("PILOT", "MAIN", "ABLATION", "DIAGNOSTIC", "HPARAM"):
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        if name == "DIAGNOSTIC":
            sheet.append(["expected", "failed", "D", str(result_path), "", ""])
    workbook.save(registry)
    xlsx = tmp_path / "out.xlsx"
    render_reports(registry, xlsx, tmp_path / "out.md")
    rows = list(load_workbook(xlsx, read_only=True)["Runs"].values)
    assert rows[1][4] == "run_id_mismatch"
