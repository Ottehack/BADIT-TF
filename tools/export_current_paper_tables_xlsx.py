#!/usr/bin/env python3
"""Export the currently auditable BADIT-TF paper-table results to Excel.

The workbook is intentionally separate from ``experiment_registry.xlsx``.  It
contains only values derivable from frozen local artifacts.  Missing values are
rendered as ``[RESULT NEEDED]`` rather than zero, blank, or paper-era numbers.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import statistics
import subprocess
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NEEDED = "[RESULT NEEDED]"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREEN = "E2F0D9"
LIGHT_YELLOW = "FFF2CC"
LIGHT_RED = "FCE4D6"
GREY = "E7E6E6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="D9E1F2")


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def git_commit(root: Path) -> str:
    return subprocess.check_output(
        ["git", "rev-parse", "HEAD"], cwd=root, text=True
    ).strip()


def stat(values: Iterable[float]) -> tuple[float, float, int]:
    xs = list(values)
    return statistics.mean(xs), statistics.stdev(xs) if len(xs) > 1 else 0.0, len(xs)


def metric3(metric: dict[str, Any], std_key: str = "std") -> list[Any]:
    return [metric["mean"], metric.get(std_key, metric.get("sample_std")), metric["n"]]


def new_sheet(wb: Workbook, title: str, description: str) -> Any:
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H1")
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 25
    ws.merge_cells("A2:H2")
    ws["A2"] = description
    ws["A2"].font = Font(italic=True, color="595959")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 32
    return ws


def add_table(ws: Any, headers: list[str], rows: list[list[Any]], start_row: int = 4) -> int:
    for col, header in enumerate(headers, 1):
        cell = ws.cell(start_row, col, header)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=THIN)
    for r_idx, row in enumerate(rows, start_row + 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(r_idx, c_idx, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN)
            if isinstance(value, float):
                cell.number_format = "0.000000"
            if value == NEEDED:
                cell.fill = PatternFill("solid", fgColor=LIGHT_RED)
                cell.font = Font(color="9C0006")
            elif isinstance(value, str) and value.lower() == "complete":
                cell.fill = PatternFill("solid", fgColor=LIGHT_GREEN)
            elif isinstance(value, str) and value.lower() == "partial":
                cell.fill = PatternFill("solid", fgColor=LIGHT_YELLOW)
        if r_idx % 2 == 0:
            for cell in ws[r_idx]:
                if cell.fill.fill_type is None:
                    cell.fill = PatternFill("solid", fgColor="F7FAFC")
    end_row = start_row + len(rows)
    ws.freeze_panes = ws.cell(start_row + 1, 1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{end_row}"
    for c_idx, header in enumerate(headers, 1):
        values = [str(header)] + [str(row[c_idx - 1]) if c_idx <= len(row) else "" for row in rows]
        width = min(max(max(map(len, values)) + 2, 10), 48)
        ws.column_dimensions[get_column_letter(c_idx)].width = width
    return end_row


def source_rel(root: Path, path: Path) -> str:
    return str(path.relative_to(root))


def make_workbook(root: Path, output: Path) -> dict[str, Any]:
    generated_at = datetime.now(timezone.utc).isoformat()
    commit = git_commit(root)
    sources: dict[str, Path] = {
        "M0_M1": root / "analysis/tables/m0_m1_official_aggregate.json",
        "M2": root / "experiments/raw_results/m2_fidelity_six_model_aggregate/result.json",
        "M3": root / "analysis/tables/m3_deployment_scale/result.json",
        "M4": root / "analysis/tables/m4_decomposition_fidelity/result.json",
        "M5": root / "analysis/tables/m5_tf_ablation_contiguous/result.json",
        "M6_manifest": root / "experiments/materials/m6_discovery_transfer_release_manifest.json",
        "M7": root / "experiments/raw_results/m7_partition_stability_small_models_formal_r1/result.json",
        "M8": root / "analysis/tables/m8_ability_evidence/result.json",
        "M9": root / "analysis/tables/m9_efficiency_partial.json",
        "A0": root / "analysis/tables/a0_calibration_config/result.json",
        "A1": root / "analysis/tables/a1_solver/result.json",
        "Spec": root / "experiments/source/BADIT_TF_EXPERIMENT_IMPLEMENTATION_SUMMARY.md",
    }
    missing_sources = [str(path) for path in sources.values() if not path.is_file()]
    if missing_sources:
        raise FileNotFoundError(f"Required sources are missing: {missing_sources}")
    data = {name: read_json(path) for name, path in sources.items() if path.suffix == ".json"}

    wb = Workbook()
    wb.remove(wb.active)

    # README
    ws = new_sheet(
        wb,
        "README",
        "当前本地可审计结果快照；按论文结果表分 sheet。它不是最终定稿表，缺失项均显式保留。",
    )
    readme_rows = [
        ["Workbook", output.name],
        ["Generated at (UTC)", generated_at],
        ["Git commit at export", commit],
        ["Evidence policy", "仅使用当前项目本地 JSON/manifest；不使用 TPAMI 旧数字补齐新实验。"],
        ["Missing-value policy", f"所有尚无正式证据的值写为 {NEEDED}，并给出原因。"],
        ["Uncertainty", "主结果 mean ± sample standard deviation；bootstrap CI 仅在已有冻结产物时填写。"],
        ["M6 heldout definition", "每 fold 对冻结的 3 个 heldout task 的 Rouge-L 做简单平均，再对 5 folds 求 mean/std。"],
        ["Important", "M5/M6/M9 为部分结果；A2 未形成正式结果。Completion_Status 给出逐表完成度。"],
    ]
    add_table(ws, ["Field", "Value"], readme_rows)

    completion = [
        ["result", "M0", "complete", "6/6 model rows", "Mixed Forward 指标不适用于当前 mixed 聚合，保持缺失。"],
        ["tab_tf_paired", "M1", "partial", "6/6 paired model rows", "现有官方聚合无层级 bootstrap CI；其余配对值和 wins 可填。"],
        ["tab_fidelity", "M2", "complete", "5/5 method rows", "eta=0.0001（锁定后的 protocol amendment）。"],
        ["tab_deployment_scale", "M3", "complete", "4/4 eta rows", "六模型 macro。"],
        ["tab_decomp_fidelity", "M4", "complete", "5/5 method rows", "eta=0.0001（protocol amendment）。"],
        ["tab_tf_ablation", "M5", "partial", "TF + contiguous", "其余四个 ablation 尚无可解释正式结果。"],
        ["tab_discovery_transfer", "M6", "partial", "Mixed 4-category/random12 complete", "Sequential 仍在执行；All15/GG heldout fold 未聚合。"],
        ["tab_partition_stability", "M7", "complete", "3/3 method rows", "三个小模型 macro。"],
        ["tab_ability_evidence", "M8", "complete", "3 models + macro", "15 cells / 108000 units。"],
        ["time", "M9", "partial", "24/55 time pairs", "绝对耗时、calibration/assignment/throughput/memory 未齐。"],
        ["tab_tf_efficiency", "M9", "partial", "time 24/55; params 25/55", "禁止对缺失 cell 插补。"],
        ["tab_tf_calib_config", "A0", "complete", "6/6 models", "大模型继承家族超参数。"],
        ["tab_tf_solver", "A1", "complete", "3 calibrated families + macro", "大模型继承家族配置。"],
        ["tab_tf_sensitivity", "A2", "result needed", "0/16 sweep points aggregated", "H0 相关证据不能替代正式 A2 单因素敏感性表。"],
    ]
    ws = new_sheet(wb, "Completion_Status", "论文结果表的当前可填范围和缺口。")
    add_table(ws, ["Sheet", "Experiment", "Status", "Fillable", "Notes"], completion)

    # M0 main table.
    m01 = data["M0_M1"]
    models = ["Gemma2-2B", "Gemma2-9B", "Llama3-3B", "Llama3-8B", "Qwen3-4B", "Qwen3-8B"]
    rows = []
    for model in models:
        mixed = m01["cells"][f"{model}.mixed.tf"]["aggregate"]["macro_rouge_l"]
        seq = m01["cells"][f"{model}.sequential.tf"]["aggregate"]
        row = [model, *metric3(mixed), NEEDED, NEEDED, NEEDED]
        for key in ("continual_score", "forget_rate", "forward", "backward"):
            row.extend(metric3(seq[key]))
        row.extend(["complete", source_rel(root, sources["M0_M1"])])
        rows.append(row)
    headers = ["Model", "Mixed Rouge-L mean", "Mixed Rouge-L std", "Mixed n", "Mixed Forward mean", "Mixed Forward std", "Mixed Forward n"]
    for label in ("Sequential Rouge-L", "Sequential Forget", "Sequential Forward", "Sequential Backward"):
        headers.extend([f"{label} mean", f"{label} std", f"{label} n"])
    headers.extend(["Status", "Source"])
    ws = new_sheet(wb, "result", "M0：TF 六模型主结果，5 seeds；数值为 mean/sample std。")
    add_table(ws, headers, rows)

    # M1 matched TF/GG table.
    rows = []
    for model in models:
        mt_tf = m01["cells"][f"{model}.mixed.tf"]["aggregate"]["macro_rouge_l"]
        mt_gg = m01["cells"][f"{model}.mixed.gg"]["aggregate"]["macro_rouge_l"]
        mt_d = m01["paired"][f"{model}.mixed"]["tf_minus_gg"]["macro_rouge_l"]
        st_tf = m01["cells"][f"{model}.sequential.tf"]["aggregate"]
        st_gg = m01["cells"][f"{model}.sequential.gg"]["aggregate"]
        st_d = m01["paired"][f"{model}.sequential"]["tf_minus_gg"]
        rows.append([
            model, mt_tf["mean"], mt_gg["mean"], mt_d["mean"], mt_d["std"], mt_d["n"],
            sum(v > 0 for v in mt_d["values"]), NEEDED, NEEDED,
            st_tf["continual_score"]["mean"], st_gg["continual_score"]["mean"],
            st_d["continual_score"]["mean"], st_d["continual_score"]["std"], st_d["continual_score"]["n"],
            sum(v > 0 for v in st_d["continual_score"]["values"]), NEEDED, NEEDED,
            st_d["forget_rate"]["mean"], st_d["forward"]["mean"], st_d["backward"]["mean"],
            "partial", source_rel(root, sources["M0_M1"]),
        ])
    ws = new_sheet(wb, "tab_tf_paired", "M1：matched TF vs GG；Delta=TF-GG，Forget 的负值更优。")
    add_table(ws, [
        "Model", "TF Mixed", "GG Mixed", "Mixed Delta mean", "Mixed Delta std", "Mixed n", "Mixed wins/5",
        "Mixed bootstrap CI lo", "Mixed bootstrap CI hi", "TF Sequential", "GG Sequential",
        "Sequential Delta mean", "Sequential Delta std", "Sequential n", "Sequential wins/5",
        "Sequential bootstrap CI lo", "Sequential bootstrap CI hi", "Forget Delta", "Forward Delta", "Backward Delta",
        "Status", "Source",
    ], rows)

    # M2 local fidelity.
    m2 = data["M2"]
    rows = []
    for method in m2["method_order"]:
        x = m2["metrics"]["macro"][method]
        rows.append([
            method, m2["eta"], x["predicted_regret_mean"], x["observed_gap_mean"],
            x["observed_gap_median_abs_macro_mean"], x["spearman"], *x["task_bootstrap_95ci"],
            x["rank_accuracy"], x["top_bottom_observed_gap_separation"], x["models"], "complete",
            source_rel(root, sources["M2"]),
        ])
    ws = new_sheet(wb, "tab_fidelity", "M2：锁定 local eta 上的六模型 fidelity macro。")
    add_table(ws, ["Method", "Eta", "Predicted regret", "Observed gap", "Median |observed gap|", "Spearman", "Bootstrap 95% CI lo", "Bootstrap 95% CI hi", "Rank accuracy", "Top-bottom separation", "Models", "Status", "Source"], rows)

    # M3 deployment scale.
    rows = []
    for eta in data["M3"]["eta_grid"]:
        x = data["M3"]["macro_by_eta"][str(eta)]
        row = [eta]
        for key in ("predicted_delta_loss_mean", "observed_delta_loss_mean", "spearman", "relative_error_median", "relative_error_p95", "max_abs_scaled_displacement"):
            row.extend([x[key]["mean"], x[key]["sample_std"]])
        row.extend([x["units"], "complete", source_rel(root, sources["M3"])])
        rows.append(row)
    headers = ["Eta"]
    for label in ("Predicted delta loss", "Observed delta loss", "Spearman", "Median relative error", "P95 relative error", "Max |scaled displacement|"):
        headers.extend([f"{label} mean", f"{label} std"])
    headers.extend(["Units", "Status", "Source"])
    ws = new_sheet(wb, "tab_deployment_scale", "M3：从局部有效区向 deployment scale 扩展的六模型曲线。")
    add_table(ws, headers, rows)

    # M4 decomposition fidelity.
    rows = []
    for method in ("tf", "contiguous", "random_balanced", "gg_dog", "raw_q"):
        x = data["M4"]["macro"][method]
        row = [method, 1e-4]
        for key in ("grouping_mean", "routing_mean", "total_pred_mean", "observed_gap_mean", "spearman_total_pred_vs_observed"):
            row.extend([x[key]["mean"], x[key]["sample_std"]])
        row.extend([x["grouping_mean"]["n_cells"], "complete", source_rel(root, sources["M4"])])
        rows.append(row)
    headers = ["Method", "Eta"]
    for label in ("Grouping pred", "Routing pred", "Total pred", "Observed gap", "Spearman"):
        headers.extend([f"{label} mean", f"{label} std"])
    headers.extend(["Cells", "Status", "Source"])
    ws = new_sheet(wb, "tab_decomp_fidelity", "M4：grouping/routing decomposition fidelity，30 cells。")
    add_table(ws, headers, rows)

    # M5 ablation: only TF and contiguous are currently valid.
    m5 = data["M5"]["macro"]
    rows = []
    for variant, status, note in [
        ("Full TF", "complete", "Reference rows from the same 3-model/3-seed comparison."),
        ("Contiguous SVD", "complete", "18/18 formal cells complete."),
        ("w/o task balancing", "result needed", "Equal probes make this mathematically identical; no distinct formal cell."),
        ("w/o curvature", "result needed", "Current raw-q evidence changes global scale; not a clean curvature-only ablation."),
        ("Raw-q", "result needed", "No matched downstream SFT ablation table."),
        ("w/o equal capacity", "result needed", "Ragged-capacity runtime not implemented as a formal comparable run."),
    ]:
        if variant == "Full TF":
            mt = m5["mixed"]["tf"]["macro_rouge_l"]
            sq = m5["sequential"]["tf"]
            vals = [mt["mean"], mt["sample_std"], sq["continual_score"]["mean"], sq["continual_score"]["sample_std"], sq["forget_rate"]["mean"], sq["forward"]["mean"], sq["backward"]["mean"]]
        elif variant == "Contiguous SVD":
            mt = m5["mixed"]["contiguous"]["macro_rouge_l"]
            sq = m5["sequential"]["contiguous"]
            vals = [mt["mean"], mt["sample_std"], sq["continual_score"]["mean"], sq["continual_score"]["sample_std"], sq["forget_rate"]["mean"], sq["forward"]["mean"], sq["backward"]["mean"]]
        else:
            vals = [NEEDED] * 7
        rows.append([variant, *vals, status, note, source_rel(root, sources["M5"])])
    ws = new_sheet(wb, "tab_tf_ablation", "M5：当前只有 Full TF 与 Contiguous SVD 可直接填表。")
    add_table(ws, ["Variant", "Mixed Rouge-L mean", "Mixed Rouge-L std", "Sequential Rouge-L mean", "Sequential Rouge-L std", "Forget", "Forward", "Backward", "Status", "Notes", "Source"], rows)

    # M6 discovery transfer, deriving the mixed folds from frozen results.
    manifest = data["M6_manifest"]
    grouped: dict[str, list[dict[str, Any]]] = {"four_category": [], "random12": []}
    m6_source_paths: list[Path] = []
    fold_rows = []
    for record in manifest["records"]:
        if record["setting"] != "mixed":
            continue
        result_path = root / "experiments/raw_results" / record["run_id"] / "result.json"
        if not result_path.is_file():
            continue
        result = read_json(result_path)
        rouge = result["metrics"]["rouge"]
        heldout = statistics.mean(rouge["per_task"][task]["rougeL"] for task in record["heldout_tasks"])
        item = {"overall": rouge["macro"]["rougeL"], "heldout": heldout, "fold": record["fold"], "run_id": record["run_id"], "path": result_path}
        grouped[record["scope"]].append(item)
        m6_source_paths.append(result_path)
        fold_rows.append([record["scope"], record["fold"], record["heldout_label"], item["overall"], heldout, record["run_id"], source_rel(root, result_path)])
    q_tf_mixed = m01["cells"]["Qwen3-4B.mixed.tf"]["aggregate"]["macro_rouge_l"]
    q_gg_mixed = m01["cells"]["Qwen3-4B.mixed.gg"]["aggregate"]["macro_rouge_l"]
    q_tf_seq = m01["cells"]["Qwen3-4B.sequential.tf"]["aggregate"]
    q_gg_seq = m01["cells"]["Qwen3-4B.sequential.gg"]["aggregate"]
    rows = []
    for label, mixed, seq, status, note in [
        ("All15 / TF", q_tf_mixed, q_tf_seq, "partial", "M0/M1 reuse; heldout-fold metric unavailable."),
        ("Four-category / TF", grouped["four_category"], None, "partial", "Mixed 5/5 complete; sequential pending."),
        ("Random12 / TF", grouped["random12"], None, "partial", "Mixed 5/5 complete; sequential pending."),
        ("GG", q_gg_mixed, q_gg_seq, "partial", "M0/M1 reuse; heldout-fold metric unavailable."),
    ]:
        if isinstance(mixed, list):
            om, osd, on = stat(x["overall"] for x in mixed)
            hm, hsd, hn = stat(x["heldout"] for x in mixed)
        else:
            om, osd, on = metric3(mixed)
            hm = hsd = hn = NEEDED
        if seq is None:
            seqvals = [NEEDED] * 5
        else:
            seqvals = [seq["continual_score"]["mean"], seq["continual_score"]["std"], seq["forget_rate"]["mean"], seq["forward"]["mean"], seq["backward"]["mean"]]
        rows.append([label, om, osd, on, hm, hsd, hn, *seqvals, NEEDED, NEEDED, status, note, source_rel(root, sources["M6_manifest"])])
    ws = new_sheet(wb, "tab_discovery_transfer", "M6：discovery scope transfer；aggregate 在上，mixed fold 明细在下。")
    end = add_table(ws, ["Discovery condition", "Mixed overall mean", "Mixed overall std", "Mixed n", "Mixed heldout mean", "Mixed heldout std", "Mixed heldout n", "Sequential Rouge-L mean", "Sequential Rouge-L std", "Sequential Forget", "Sequential Forward", "Sequential Backward", "Discovery regret", "Fidelity", "Status", "Notes", "Source"], rows)
    add_table(ws, ["Scope", "Fold", "Heldout category", "Overall Rouge-L", "Heldout Rouge-L", "Run ID", "Source"], sorted(fold_rows), start_row=end + 3)

    # M7 partition stability.
    rows = []
    for method in ("tf", "random_balanced", "contiguous"):
        x = data["M7"]["macro"][method]
        row = [method]
        for key in ("task_split_ari", "seed_ari", "batch_ari", "cross_layer_cosine"):
            row.extend([x[key]["mean"], x[key]["sample_std"], x[key]["n"]])
        row.extend(["complete", source_rel(root, sources["M7"])])
        rows.append(row)
    headers = ["Method"]
    for label in ("Task-split ARI", "Seed ARI", "Batch ARI", "Cross-layer cosine"):
        headers.extend([f"{label} mean", f"{label} std", f"{label} n"])
    headers.extend(["Status", "Source"])
    ws = new_sheet(wb, "tab_partition_stability", "M7：三个小模型的 partition stability macro。")
    add_table(ws, headers, rows)

    # M8 intervention / ability evidence.
    rows = []
    for model in ("Gemma2-2B", "Llama3-3B", "Qwen3-4B"):
        x = data["M8"]["per_model"][model]
        metrics = x["metrics"]
        row = [model]
        for key in ("target_deletion_increase", "control_deletion_increase", "specificity_gap", "sufficiency", "composition_gain", "effective_experts", "top1_mass"):
            row.extend([metrics[key]["mean"], metrics[key]["sample_std"]])
        row.extend([metrics["specificity_gap"]["positive_count"], metrics["composition_gain"]["positive_count"], x["units"], "complete", source_rel(root, sources["M8"])])
        rows.append(row)
    macro = data["M8"]["macro"]
    row = ["Macro"]
    for key in ("target_deletion_increase", "control_deletion_increase", "specificity_gap", "sufficiency", "composition_gain", "effective_experts", "top1_mass"):
        row.extend([macro[key]["mean"], macro[key]["sample_std"]])
    row.extend([macro["specificity_gap"]["positive_count"], macro["composition_gain"]["positive_count"], data["M8"]["total_units"], "complete", source_rel(root, sources["M8"])])
    rows.append(row)
    headers = ["Model"]
    for label in ("Target deletion increase", "Control deletion increase", "Specificity gap", "Sufficiency", "Composition gain", "Effective experts", "Top1 mass"):
        headers.extend([f"{label} mean", f"{label} std"])
    headers.extend(["Specificity positive count", "Composition positive count", "Units", "Status", "Source"])
    ws = new_sheet(wb, "tab_ability_evidence", "M8：deletion/composition intervention 证据；3 models × 5 seeds。")
    add_table(ws, headers, rows)

    # M9 timing details and efficiency macro.
    m9 = data["M9"]
    rows = []
    for cell, values in sorted(m9["by_model_setting"].items()):
        t = values["train_time_tf_over_gg"]
        p = values["trainable_params_tf_over_gg"]
        rows.append([cell, t["mean"] if t["n"] else NEEDED, t["sample_std"] if t["n"] else NEEDED, t["n"], p["mean"] if p["n"] else NEEDED, p["sample_std"] if p["n"] else NEEDED, p["n"], NEEDED, NEEDED, "partial", source_rel(root, sources["M9"])])
    ws = new_sheet(wb, "time", "M9：当前可恢复的逐 model/setting 相对训练时间；绝对 timing 尚未形成完整表。")
    add_table(ws, ["Model.setting", "Train time TF/GG mean", "Train time TF/GG std", "Time n", "Trainable params TF/GG mean", "Params std", "Params n", "Absolute TF time", "Absolute GG time", "Status", "Source"], rows)
    rows = []
    coverage = m9["coverage"]
    for metric, key in [
        ("Train time TF/GG", "train_time_tf_over_gg"),
        ("Trainable params TF/GG", "trainable_params_tf_over_gg"),
    ]:
        x = m9["available_macro"][key]
        rows.append([metric, x["mean"], x["sample_std"], x["n"], coverage["required_paired_cells"], x["n"] / coverage["required_paired_cells"], "partial", source_rel(root, sources["M9"])])
    for metric, covkey in [
        ("Calibration GPU hours", "calibration_gpu_hours_cells"),
        ("Assignment CPU seconds", "assignment_cpu_seconds_cells"),
        ("Peak calibration memory", "peak_calibration_memory_cells"),
        ("Inference throughput TF/GG", "inference_throughput_cells"),
    ]:
        rows.append([metric, NEEDED, NEEDED, coverage[covkey], coverage["required_paired_cells"], 0.0, "result needed", source_rel(root, sources["M9"])])
    ws = new_sheet(wb, "tab_tf_efficiency", "M9：只报告已恢复覆盖率；不对缺失的 efficiency cell 插补。")
    add_table(ws, ["Metric", "Mean", "Std", "Available cells", "Required cells", "Coverage", "Status", "Source"], rows)

    # A0 calibration config.
    rows = []
    for x in data["A0"]["rows"]:
        rows.append([x["model"], x["assignment_probes_per_task"], x["fisher_probes_per_task"], x["fidelity_probes_per_task"], x["epsilon_f"], x["solver_restarts"], x["max_iterations"], "complete", x["selection_source"], x["selection_source_sha256"]])
    ws = new_sheet(wb, "tab_tf_calib_config", "A0：逐模型实际 calibration/solver 配置。")
    add_table(ws, ["Model", "Assignment probes/task", "Fisher probes/task", "Fidelity probes/task", "epsilon_f", "Restarts", "Max iterations", "Status", "Selection source", "Selection SHA256"], rows)

    # A1 solver diagnostics.
    rows = []
    for x in data["A1"]["rows"]:
        rows.append([x["model"], x["selected_trial_id"], x["iterations_mean"], x["convergence_rate"], x["best_regret_mean"], x["restart_spread_cv_mean"], x["layers"], "complete", x["solver_audit_path"], x["solver_audit_sha256"]])
    x = data["A1"]["macro"]
    rows.append(["Macro", "", x["iterations_mean"], x["convergence_rate"], x["best_regret_mean"], x["restart_spread_cv_mean"], sum(r["layers"] for r in data["A1"]["rows"]), "complete", source_rel(root, sources["A1"]), sha256(sources["A1"])])
    ws = new_sheet(wb, "tab_tf_solver", "A1：三个模型家族 calibration lock 的 solver 收敛与 restart 稳定性。")
    add_table(ws, ["Model", "Selected trial", "Iterations mean", "Convergence rate", "Best regret mean", "Restart spread CV", "Layers", "Status", "Audit source", "Audit SHA256"], rows)

    # A2 remains explicitly missing.
    rows = []
    for axis, values in [
        ("assignment probes/task", [4, 8, 16, 32]),
        ("Fisher samples/task", [4, 8, 16, 32]),
        ("epsilon_f", [1e-4, 1e-3, 1e-2, 1e-1]),
        ("restarts", [1, 3, 5, 10]),
    ]:
        for value in values:
            rows.append(["Qwen3-4B", axis, value, NEEDED, NEEDED, NEEDED, "result needed", "No formal A2 one-factor aggregate is available."])
    ws = new_sheet(wb, "tab_tf_sensitivity", "A2：固定预注册 grid；当前不以 H0 证据冒充正式 sensitivity 结果。")
    add_table(ws, ["Model", "Axis", "Value", "Stability ARI", "Fidelity", "Validation metric", "Status", "Notes"], rows)

    # Artifact index, including all M6 mixed result files used.
    artifact_rows = []
    all_sources = list(sources.items()) + [(f"M6_{p.parent.name}", p) for p in sorted(set(m6_source_paths))]
    for label, path in all_sources:
        artifact_rows.append([label, source_rel(root, path), path.stat().st_size, sha256(path), "used"])
    ws = new_sheet(wb, "Artifact_Index", "本工作簿直接读取的 artifact 及 SHA256。")
    add_table(ws, ["Label", "Relative path", "Bytes", "SHA256", "Role"], artifact_rows)

    # Workbook-wide polish and explicit missing marker checks.
    for ws in wb.worksheets:
        ws.auto_filter.ref = ws.auto_filter.ref or ws.dimensions
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.outlinePr.summaryBelow = True
        ws.freeze_panes = ws.freeze_panes or "A5"

    output.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(prefix=output.stem + ".", suffix=".xlsx", dir=output.parent)
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        wb.save(tmp)
        os.replace(tmp, output)
    finally:
        if tmp.exists():
            tmp.unlink()

    # Re-open to ensure the ZIP and every sheet are readable.
    check = load_workbook(output, data_only=False, read_only=True)
    sheet_summary = {ws.title: {"rows": ws.max_row, "columns": ws.max_column} for ws in check.worksheets}
    check.close()
    return {
        "schema_version": 1,
        "generated_at_utc": generated_at,
        "git_commit": commit,
        "workbook": source_rel(root, output),
        "workbook_sha256": sha256(output),
        "workbook_bytes": output.stat().st_size,
        "sheets": sheet_summary,
        "sources": {name: {"path": source_rel(root, path), "sha256": sha256(path)} for name, path in all_sources},
        "missing_marker": NEEDED,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[2])
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    root = args.root.resolve()
    output = args.output.resolve() if args.output else root / "experiments/results/BADIT_TF_CURRENT_TABLE_RESULTS.xlsx"
    manifest = make_workbook(root, output)
    manifest_path = output.with_suffix(".manifest.json")
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"workbook": str(output), "manifest": str(manifest_path), "sha256": manifest["workbook_sha256"], "sheets": len(manifest["sheets"])}, ensure_ascii=False))


if __name__ == "__main__":
    main()
