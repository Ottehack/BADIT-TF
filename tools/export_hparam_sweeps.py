#!/usr/bin/env python3
"""Export every registered hyperparameter row plus frozen H0-H4 locks."""
from __future__ import annotations
import argparse,hashlib,json,os,tempfile
from collections import Counter
from datetime import datetime,timezone
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
ROOT=Path(__file__).resolve().parents[2]
LOCKS=[
'experiments/raw_results/h0_calibration_sweep_qwen3_4b_seed24001_remote_l20z/h0_selection_lock.json','experiments/raw_results/h0_calibration_sweep_llama3_3b_seed24001_local_a100/h0_selection_lock.json','experiments/raw_results/h0_calibration_sweep_gemma2_2b_seed24001_local_a100/h0_selection_lock.json','experiments/raw_results/h1_qwen3_4b_one_factor_aggregate/selection_lock.json','experiments/raw_results/h1_llama3_3b_one_factor_aggregate/selection_lock.json','experiments/raw_results/h1_gemma2_2b_one_factor_aggregate/selection_lock.json','experiments/raw_results/h2_joint_search_aggregate/selection_lock.json','experiments/raw_results/h3_confirmation_aggregate/selection_lock.json','experiments/raw_results/h3_qwen_mixed_timing_replay_aggregate/selection_lock.json','experiments/raw_results/h4_confirmation_aggregate/selection_lock.json']
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def style(s):
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1F4E78');c.alignment=Alignment(wrap_text=True)
 for col in s.columns:
  letter=col[0].column_letter;s.column_dimensions[letter].width=min(max(max(len(str(c.value or '')) for c in col)+2,10),50)
def main():
 p=argparse.ArgumentParser();p.add_argument('--registry',type=Path,default=ROOT/'experiments/experiment_registry.xlsx');p.add_argument('--output',type=Path,default=ROOT/'experiments/results/BADIT_TF_HPARAM_SWEEPS.xlsx');a=p.parse_args()
 source=load_workbook(a.registry,read_only=True,data_only=False);hp=source['HPARAM'];headers=[c.value for c in hp[1]];rows=[list(r) for r in hp.iter_rows(min_row=2,values_only=True) if r[0]];idx={v:i for i,v in enumerate(headers)}
 formal=[r for r in rows if str(r[idx['todo_id']]).split('-')[0] in {'H0','H1','H2','H3','H4'} or str(r[idx['setting_group']]).startswith(('H0','H1','H2','H3','H4'))]
 wb=Workbook();readme=wb.active;readme.title='README';readme.append(['Field','Value']);readme.append(['Policy','Full HPARAM registry snapshot plus non-destructive H0-H4 formal view; failed/blocked/superseded rows retained.']);readme.append(['Registry SHA256',sha(a.registry)]);readme.append(['All HPARAM rows',len(rows)]);readme.append(['Formal H0-H4 rows',len(formal)]);readme.append(['Generated UTC',datetime.now(timezone.utc).isoformat()]);style(readme)
 for title,data in [('All_HPARAM_Rows',rows),('Formal_H0_H4',formal)]:
  s=wb.create_sheet(title);s.append(headers)
  for r in data:s.append(r)
  style(s)
 locks=wb.create_sheet('Selection_Locks');locks.append(['Path','Bytes','SHA256','Top-level keys'])
 lock_rows=[]
 for rel in LOCKS:
  path=ROOT/rel
  if not path.is_file():raise FileNotFoundError(path)
  payload=json.loads(path.read_text());row=[rel,path.stat().st_size,sha(path),', '.join(sorted(payload)[:20])];locks.append(row);lock_rows.append({'path':rel,'bytes':path.stat().st_size,'sha256':sha(path)})
 style(locks);source.close();a.output.parent.mkdir(parents=True,exist_ok=True);fd,tmpname=tempfile.mkstemp(prefix=a.output.stem+'.',suffix='.xlsx',dir=a.output.parent);os.close(fd);tmp=Path(tmpname)
 try:wb.save(tmp);os.replace(tmp,a.output)
 finally:
  if tmp.exists():tmp.unlink()
 status=Counter(str(r[idx['status']]) for r in rows);todo=Counter(str(r[idx['todo_id']]) for r in formal)
 audit={'schema_version':1,'status':'complete','generated_at_utc':datetime.now(timezone.utc).isoformat(),'registry':str(a.registry.relative_to(ROOT)),'registry_sha256':sha(a.registry),'output':str(a.output.relative_to(ROOT)),'output_sha256':sha(a.output),'all_hparam_rows':len(rows),'formal_h0_h4_rows':len(formal),'status_counts':dict(status),'formal_todo_counts':dict(todo),'selection_locks':lock_rows,'assertions':{'all_rows_preserved':len(rows)==hp.max_row-1,'failed_rows_retained':status.get('failed',0)>0,'blocked_rows_retained':status.get('blocked',0)>0,'ten_locks_indexed':len(lock_rows)==10,'required_columns_present':all(k in idx for k in ('run_id','status','config_path','val_data','test_data','metrics'))}}
 ap=a.output.with_suffix('.audit.json');ap.write_text(json.dumps(audit,indent=2,sort_keys=True)+'\n');print(json.dumps({'output':str(a.output),'sha256':sha(a.output),'rows':len(rows),'formal':len(formal)},sort_keys=True))
if __name__=='__main__':main()
