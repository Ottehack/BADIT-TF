#!/usr/bin/env python3
"""Build a compact, result-first workbook for human TPAMI run selection."""

from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "experiments/results/BADIT_TF_TPAMI_ALL_RUNS.xlsx"
OUTPUT = ROOT / "experiments/results/BADIT_TF_TPAMI_RUN_CANDIDATES.xlsx"
MANIFEST = OUTPUT.with_suffix(".manifest.json")
TABLES = ("I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "XI", "XII", "XIX", "XX", "XXI")
IDENTITY = ("backbone", "method", "seed", "run_id", "included_in_current_summary")

# Only scalar fields that can plausibly populate or help choose a row in the
# corresponding paper table. Management, provenance and diagnostic internals
# deliberately stay in the full all-runs workbook.
TOKENS = {
    "I": ("rouge", "forward", "continual_score", "forget_rate", "backward"),
    "II": ("rouge", "forward", "continual_score", "forget_rate", "backward"),
    "III": ("predicted_regret", "spearman", "rank_accuracy", "observed_gap", "task_bootstrap_95ci", "top_bottom"),
    "IV": ("predicted_regret", "spearman", "rank_accuracy", "observed_gap", "top_bottom"),
    "V": ("grouping", "routing", "total", "observed_gap", "spearman", "regret"),
    "VI": ("rouge", "continual_score", "forget_rate", "predicted_regret", "spearman"),
    "VII": ("rouge", "forward", "forget_rate", "seen", "unseen", "overall"),
    "VIII": ("ari", "cosine", "stability"),
    "IX": ("target_deletion", "control_deletion", "specificity_gap", "sufficiency", "top1_score", "top4_score", "composition_gain"),
    "XI": ("ratio", "wall", "seconds", "gpu_hours", "cost"),
    "XII": ("parameter", "memory", "throughput", "seconds", "gpu_hours", "ratio"),
    "XIX": ("rouge", "forward", "continual_score", "forget_rate", "backward", "regret", "stability", "fidelity"),
    "XX": ("iteration", "convergence", "regret", "restart", "objective"),
    "XXI": ("regret", "stability", "fidelity", "spearman", "noise", "time", "seconds"),
}


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def short_metric_name(value: str) -> str:
    value = value.removeprefix("metric.").replace("metrics.", "")
    value = value.replace("macro.", "").replace("continual.", "")
    return value


def useful(column: str, table: str) -> bool:
    lowered = column.lower()
    if not column.startswith("metric."):
        return False
    if any(part in lowered for part in ("assertion", "sha256", "path", "count", "parameters.changed", "rank_model", ".models.", ".per_model.")):
        return False
    return any(token in lowered for token in TOKENS[table])


def has_numeric_result(row: tuple, indices: list[int]) -> bool:
    return any(isinstance(row[index], (int, float)) and not isinstance(row[index], bool) for index in indices)


def main() -> None:
    source = load_workbook(SOURCE, read_only=True, data_only=True)
    output = load_workbook(SOURCE)
    # Remove the verbose sheets and retain the original audited summary sheets.
    for name in list(output.sheetnames):
        if name == "All_Runs_Index" or name.endswith("_All_Runs"):
            del output[name]
    counts = {}
    for table in TABLES:
        source_sheet = source[f"Table_{table}_All_Runs"]
        headers = [str(cell.value or "") for cell in source_sheet[4]]
        positions = {name: index for index, name in enumerate(headers)}
        metric_columns = [name for name in headers if useful(name, table)]
        # Prefer direct run metrics over nested aggregate replicas and keep the
        # sheet human-sized. The full workbook remains available if needed.
        metric_columns.sort(key=lambda name: (name.count("."), len(name), name))
        metric_columns = metric_columns[:12]
        # Repeated nested aliases can make selection harder. Keep deterministic
        # unique labels while retaining the original scalar values.
        labels = []
        seen = {}
        for column in metric_columns:
            label = short_metric_name(column)
            seen[label] = seen.get(label, 0) + 1
            labels.append(label if seen[label] == 1 else f"{label} [{seen[label]}]")
        metric_indices = [positions[name] for name in metric_columns]
        rows = []
        for values in source_sheet.iter_rows(min_row=5, values_only=True):
            if str(values[positions["status"]]).lower() != "complete":
                continue
            if not values[positions["local_result_path"]]:
                continue
            if not has_numeric_result(values, metric_indices):
                continue
            rows.append(values)
        name = f"Table_{table}_Candidates"
        sheet = output.create_sheet(name)
        sheet.append([f"Table {table} — successful result candidates"])
        sheet.append(["Only complete runs with a local result and at least one table-relevant numeric metric are retained."])
        sheet.append([])
        display_headers = ["Model", "Method", "Seed", "Run ID", "Currently used", *labels]
        sheet.append(display_headers)
        for values in rows:
            sheet.append([
                values[positions[column]] for column in IDENTITY
            ] + [values[index] for index in metric_indices])
        for cell in sheet[4]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "A5"
        sheet.auto_filter.ref = f"A4:{get_column_letter(sheet.max_column)}{max(4, sheet.max_row)}"
        sheet.column_dimensions["A"].width = 18
        sheet.column_dimensions["B"].width = 30
        sheet.column_dimensions["C"].width = 10
        sheet.column_dimensions["D"].width = 54
        sheet.column_dimensions["E"].width = 14
        for index in range(6, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 19
        counts[name] = len(rows)
    source.close()
    readme = output["README"]
    readme["A16"] = "Candidate workbook"
    readme["B16"] = "Only successful result-bearing runs and table-relevant metrics are shown in *_Candidates sheets; paths and audit fields are omitted."
    temporary = OUTPUT.with_suffix(".xlsx.tmp")
    output.save(temporary)
    os.replace(temporary, OUTPUT)
    payload = {
        "schema_version": 1,
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "source": str(SOURCE.relative_to(ROOT)), "source_sha256": sha(SOURCE),
        "output": str(OUTPUT.relative_to(ROOT)), "output_sha256": sha(OUTPUT),
        "candidate_counts": counts,
        "filter": "status=complete AND local result exists AND table-relevant numeric result exists",
        "omitted": "paths, hashes, runners, hosts, logs, notes, failures, and non-table diagnostics",
    }
    MANIFEST.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n")
    print(json.dumps(payload, sort_keys=True))


if __name__ == "__main__":
    main()
