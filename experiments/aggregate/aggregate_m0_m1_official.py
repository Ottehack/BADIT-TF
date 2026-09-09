#!/usr/bin/env python3
"""Aggregate the 120 locked M0/M1 official-test runs without selecting on test.

The dispatch manifests supply the pre-registered base run IDs.  When an
original row is intentionally retained as failed, exactly one completed
descendant recovery supplies that base run's terminal metric.  The mapping is
written out, so every reported aggregate remains traceable to a registry row.
"""

from __future__ import annotations

import json
import math
import statistics
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
REGISTRY = ROOT / "experiments/experiment_registry.xlsx"
DISPATCHES = [
    ROOT / "experiments/materials/m0_m1_final_dispatch_manifest.json",
    ROOT / "experiments/materials/m0_m1_qwen3_4b_mixed_timing_extension_manifest.json",
]
OUTPUT = ROOT / "analysis/tables/m0_m1_official_aggregate.json"
MARKDOWN = ROOT / "analysis/tables/M0_M1_OFFICIAL_AGGREGATE.md"


def stat(values: list[float]) -> dict[str, object]:
    if not values:
        raise ValueError("cannot aggregate no values")
    return {
        "mean": statistics.fmean(values),
        "std": statistics.stdev(values) if len(values) > 1 else 0.0,
        "n": len(values),
    }


def metric(metrics: dict, setting: str) -> dict[str, float]:
    if setting == "mixed":
        rouge = metrics["rouge"]["macro"]
        return {
            "macro_rouge_l": float(rouge["rougeL"]),
            "macro_rouge_1": float(rouge["rouge1"]),
            "exact_match": float(rouge["exact_match"]),
        }
    continual = metrics["continual"]
    return {
        "continual_score": float(continual["continual_score"]),
        "forget_rate": float(continual["forget_rate"]),
        "forward": float(continual["forward"]),
        "backward": float(continual["backward"]),
    }


def resolved_rows() -> tuple[list[dict], dict[str, dict]]:
    book = load_workbook(REGISTRY, read_only=True, data_only=True)
    sheet = book["HPARAM"]
    columns = [cell.value for cell in sheet[1]]
    rows = [dict(zip(columns, values)) for values in sheet.iter_rows(min_row=2, values_only=True)]
    by_id = {row["run_id"]: row for row in rows if row.get("run_id")}
    trials = []
    for dispatch_path in DISPATCHES:
        dispatch = json.loads(dispatch_path.read_text(encoding="utf-8"))
        trials.extend(dispatch["trials"])
    base_ids = [trial["run_id"] for trial in trials]
    if len(base_ids) != len(set(base_ids)):
        raise ValueError("duplicate base run IDs across M0/M1 dispatch manifests")
    resolved: dict[str, dict] = {}
    for trial in trials:
        base_id = trial["run_id"]
        original = by_id.get(base_id)
        if original is None:
            raise ValueError(f"dispatch run absent from registry: {base_id}")
        candidates = []
        if original.get("status") == "complete":
            candidates.append(original)
        candidates.extend(
            row for run_id, row in by_id.items()
            if run_id.startswith(base_id + "_") and row.get("status") == "complete"
        )
        if len(candidates) != 1:
            raise ValueError(f"expected exactly one final row for {base_id}, got {[r['run_id'] for r in candidates]}")
        selected = candidates[0]
        canonical = ROOT / f"experiments/raw_results/m0_m1_official_result_records/{selected['run_id']}/result.json"
        if not canonical.is_file():
            raise FileNotFoundError(canonical)
        metrics = json.loads(canonical.read_text(encoding="utf-8"))["metrics"]
        assertions = metrics.get("assertions", {})
        if not assertions or not all(value is True for value in assertions.values()):
            raise ValueError(f"non-passing selected result: {selected['run_id']}")
        resolved[base_id] = {"trial": trial, "registry_run_id": selected["run_id"], "metrics": metrics}
    return rows, resolved


def main() -> None:
    _, resolved = resolved_rows()
    by_cell: dict[tuple[str, str, str], list[tuple[int, dict]]] = defaultdict(list)
    for base_id, record in resolved.items():
        trial = record["trial"]
        by_cell[(trial["model"], trial["setting"], trial["variant"])].append((int(trial["seed"]), record))
    cells = {}
    for (model, setting, variant), records in sorted(by_cell.items()):
        records.sort()
        per_seed = []
        for seed, record in records:
            per_seed.append({
                "seed": seed,
                "base_run_id": record["trial"]["run_id"],
                "terminal_registry_run_id": record["registry_run_id"],
                "values": metric(record["metrics"], setting),
            })
        metric_names = per_seed[0]["values"].keys()
        cells[f"{model}.{setting}.{variant}"] = {
            "model": model,
            "setting": setting,
            "variant": variant,
            "per_seed": per_seed,
            "aggregate": {name: stat([row["values"][name] for row in per_seed]) for name in metric_names},
        }
    paired = {}
    for model, setting in sorted({(v["model"], v["setting"]) for v in cells.values()}):
        tf = cells[f"{model}.{setting}.tf"]
        gg = cells[f"{model}.{setting}.gg"]
        tf_by_seed = {row["seed"]: row for row in tf["per_seed"]}
        gg_by_seed = {row["seed"]: row for row in gg["per_seed"]}
        if set(tf_by_seed) != set(gg_by_seed):
            raise ValueError(f"unpaired seeds for {model}/{setting}")
        names = tf["per_seed"][0]["values"].keys()
        differences = {
            name: [tf_by_seed[seed]["values"][name] - gg_by_seed[seed]["values"][name] for seed in sorted(tf_by_seed)]
            for name in names
        }
        paired[f"{model}.{setting}"] = {
            "seeds": sorted(tf_by_seed),
            "tf_minus_gg": {name: stat(values) | {"values": values} for name, values in differences.items()},
        }
    payload = {
        "schema_version": 1,
        "created_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
        "dispatch_manifests": [str(path.relative_to(ROOT)) for path in DISPATCHES],
        "selection_policy": "pre-registered H3/H4 validation locks; official test was never used for selection",
        "base_run_count": len(resolved),
        "cells": cells,
        "paired": paired,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    lines = ["# M0/M1 Official Aggregate", "", "All values are mean ± sample std over five pre-registered seeds.", ""]
    lines += ["## TF", "", "| Model | Setting | Primary official metric |", "|---|---|---:|"]
    for key, cell in cells.items():
        if cell["variant"] != "tf":
            continue
        primary = "macro_rouge_l" if cell["setting"] == "mixed" else "continual_score"
        value = cell["aggregate"][primary]
        lines.append(f"| {cell['model']} | {cell['setting']} | {primary} {value['mean']:.4f} ± {value['std']:.4f} |")
    lines += ["", "## Paired TF − GG", "", "| Model | Setting | Primary difference |", "|---|---|---:|"]
    for key, row in paired.items():
        model, setting = key.rsplit(".", 1)
        primary = "macro_rouge_l" if setting == "mixed" else "continual_score"
        value = row["tf_minus_gg"][primary]
        lines.append(f"| {model} | {setting} | {primary} {value['mean']:.4f} ± {value['std']:.4f} |")
    MARKDOWN.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps({"base_runs": len(resolved), "cells": len(cells), "output": str(OUTPUT.relative_to(ROOT))}, sort_keys=True))


if __name__ == "__main__":
    main()
